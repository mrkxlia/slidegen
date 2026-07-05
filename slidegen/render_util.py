"""
render_util.py — render_*.py 群で共有する実装ヘルパー。

variant 解決・複数項目テキスト描画・DSL数値パース・列ジオメトリなど、各 render_*.py に
コピペで散らばっていた同一パターンをここへ集約する。DSL/公開APIには関与しない、
レンダラ内部専用の実装ヘルパー（型登録は register/register_many のまま・public API は不変）。
"""
from __future__ import annotations
import logging

from pptx.util import Pt
from pptx.enum.text import PP_ALIGN

from .parser import Slide, split_emphasis

_log = logging.getLogger(__name__)


def variant_name(data: Slide) -> str:
    """props['variant'] があればそれを、無ければスライド型名を返す（variant 解決の第一歩）。"""
    return data.props.get("variant") or data.type


def resolve_variant(data: Slide, variants: dict, default: dict) -> dict:
    """VARIANTS 辞書から variant_name(data) の定義のコピーを返す。無ければ default のコピー。"""
    return dict(variants.get(variant_name(data), default))


def block_items(blk) -> list:
    """ブロックの本文項目リスト。lines は各行が1項目。rows は 'ラベル: 値'。"""
    parts = list(blk.lines)
    parts += [f"{lbl}: {val}" if lbl else val for lbl, val in blk.rows]
    return parts


def add_items_text(slide, x, y, w, h, theme, items, *, size, anchor, bullet):
    """複数項目を1項目=1段落で描く。bullet=Trueなら先頭に「・」。"""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Pt(2)
    tf.margin_top = tf.margin_bottom = Pt(2)
    tf.vertical_anchor = anchor
    for j, item in enumerate(items):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(4)
        prefix = "・" if bullet else ""
        for text, em in split_emphasis(prefix + item):
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.name = theme.font
            r.font.bold = em
            r.font.color.rgb = theme.rgb("accent") if em else theme.rgb("ink")
    return box


def parse_number(s: str, *, context: str = "") -> float:
    """DSL の数値文字列（"+40"/"-15"/"1,200" 等）を float にパースする。
    解釈できなければ 0.0 を返し警告ログを残す（AI 生成 DSL の typo で値が黙って 0 として
    描画される事故を可視化するため。theme.py の型付き except + logging.warning に倣う）。
    """
    try:
        return float(s.replace(",", "").replace("+", ""))
    except (ValueError, AttributeError):
        where = f"（{context}）" if context else ""
        _log.warning("数値として解釈できない値 %r%s を 0.0 として扱います。", s, where)
        return 0.0


def columns_geometry(total_w, n: int, gap):
    """n列を gap 間隔で並べたときの1列幅。`(total_w - gap*(n-1)) / n` の重複を集約する。"""
    return (total_w - gap * (n - 1)) / n if n else total_w


def fill_shape(shape, theme, color_name, *, no_shadow=False):
    """塗り潰し＋線なしのネイティブ図形にする共通処理。
    no_shadow=True の呼び出し元（render_relations.py の _add_oval/_add_triangle 等）は
    ここで shadow も無効化する。add_rect 経由の呼び出し元は従来どおり呼び出し側で
    shadow.inherit=False を設定するため、ここでは既定 no_shadow=False（挙動不変）。
    """
    shape.fill.solid()
    shape.fill.fore_color.rgb = theme.rgb(color_name)
    shape.line.fill.background()
    if no_shadow:
        shape.shadow.inherit = False
