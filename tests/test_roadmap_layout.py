"""
test_roadmap_layout.py — roadmap（render_frameworks2.render_roadmap）のバー詰めとマイルストーンの回帰テスト。

- 期間が重ならないバーは同じサブ行に並ぶ（直列の工程が細切れの行に分かれない）
- 期間が重なるバーだけ下のサブ行に回る
- milestones "期間名:ラベル" で▲とラベルが描かれ、解決できない期間名はスキップされる
"""
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE

import slidegen  # 全 render_* を登録
from slidegen.parser import parse
from slidegen.render import build
from slidegen.render_frameworks2 import _roadmap_pack

_SRC = '''slide roadmap
  headline "設立に向けたスケジュール"
  periods "5月" "6月" "7月" "8月"
  milestones "6月:設立宣言" "9月:存在しない" "8月：正式設立"
  col "分科会1"
    5月 "A"
    6月 "B"
    7月-8月 "C"
'''


def test_pack_puts_non_overlapping_bars_on_same_row():
    assert _roadmap_pack([(0, 0), (1, 1), (2, 3)]) == [0, 0, 0]


def test_pack_moves_only_overlapping_bars_down():
    assert _roadmap_pack([(0, 2), (1, 3), (3, 3)]) == [0, 1, 0]


def _shapes():
    return list(build(parse(_SRC)).slides[0].shapes)


def _text_shape(shapes, text):
    return next(s for s in shapes if s.has_text_frame and s.text_frame.text == text)


def test_sequential_bars_share_one_row():
    shapes = _shapes()
    tops = {_text_shape(shapes, t).top for t in ("A", "B", "C")}
    assert len(tops) == 1


def test_milestones_draw_triangle_and_label_and_skip_unknown_period():
    shapes = _shapes()
    tris = [s for s in shapes
            if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and s.auto_shape_type == MSO_SHAPE.ISOSCELES_TRIANGLE]
    assert len(tris) == 2
    texts = {s.text_frame.text for s in shapes if s.has_text_frame}
    assert {"設立宣言", "正式設立"} <= texts
    assert not any("存在しない" in t for t in texts)
